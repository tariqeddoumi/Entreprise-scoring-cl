#!/usr/bin/env python3
"""
Convertit la note méthodologique Markdown en document Word professionnel.

Sous-ensemble Markdown pris en charge : titres # à ####, paragraphes,
tableaux GFM, listes à puces et numérotées, blocs de code, citations,
règles horizontales, et le formatage en ligne gras / italique / code.

Usage : python3 scripts/md-to-docx.py entree.md sortie.docx "Titre" "Sous-titre"
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BRAND = RGBColor(0x0F, 0x3D, 0x63)      # bleu institutionnel
GREY = RGBColor(0x5C, 0x64, 0x70)
HEADER_FILL = "0F3D63"
BAND_FILL = "F2F5F8"


# --------------------------------------------------------------------------
# Styles
# --------------------------------------------------------------------------
def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    specs = [
        ("Heading 1", 18, True, BRAND, 20, 10),
        ("Heading 2", 14, True, BRAND, 16, 8),
        ("Heading 3", 11.5, True, RGBColor(0x1D, 0x1D, 0x1D), 12, 6),
        ("Heading 4", 10.5, True, GREY, 10, 4),
    ]
    for name, size, bold, color, before, after in specs:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.space_before = Pt(6)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.left_indent = Inches(0.25)


def shade(cell, fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


# --------------------------------------------------------------------------
# Formatage en ligne
# --------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text: str, base_bold=False, base_italic=False):
    # Liens Markdown : conserver le libellé, écarter l'URL.
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("<br>", " ").replace("\\|", "|")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x8A, 0x2B, 0x2B)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if base_bold:
            run.bold = True
        if base_italic:
            run.italic = True
    return paragraph


# --------------------------------------------------------------------------
# Page de garde et table des matières
# --------------------------------------------------------------------------
def cover_page(doc: Document, title: str, subtitle: str) -> None:
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = BRAND

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(14)
    r.font.color.rgb = GREY

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    for line in [
        "Version 2.0 — 19 août 2026",
        "Direction des Risques — Modèles de crédit entreprises",
        "Usage interne",
    ]:
        r = p.add_run(line + "\n")
        r.font.size = Pt(10.5)
        r.font.color.rgb = GREY

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Document de travail. Les pondérations, seuils et barèmes présentés ne constituent "
        "ni des seuils réglementaires, ni des probabilités de défaut calibrées, ni une "
        "validation de modèle."
    )
    r.font.size = Pt(8.5)
    r.italic = True
    r.font.color.rgb = GREY

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    r = p.add_run("Table des matières")
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = BRAND
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Ouvrir dans Word puis actualiser le champ (F9) pour générer la table."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (fld, instr, sep, placeholder, end):
        run._r.append(el)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def page_numbers(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, attr, text in (
        ("w:fldChar", "begin", None),
        ("w:instrText", None, "PAGE"),
        ("w:fldChar", "end", None),
    ):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        if text:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY


# --------------------------------------------------------------------------
# Tableaux
# --------------------------------------------------------------------------
def split_row(line: str):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    # Sépare sur | non échappé.
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def add_table(doc: Document, rows, usable_width_emu: int) -> None:
    header = rows[0]
    body = rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Largeurs proportionnelles au contenu, bornées pour rester lisibles.
    weights = []
    for i in range(ncols):
        cells = [header[i]] + [r[i] if i < len(r) else "" for r in body]
        longest = max((len(c) for c in cells), default = 1)
        weights.append(min(max(longest, 6), 60))
    total = sum(weights)
    widths = [int(usable_width_emu * w / total) for w in weights]
    table.columns_width = widths

    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        cell.width = widths[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(2)
        add_inline(para, text, base_bold=True)
        for run in para.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, HEADER_FILL)

    for idx, row in enumerate(body):
        cells = table.add_row().cells
        for i in range(ncols):
            cell = cells[i]
            cell.width = widths[i]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            add_inline(para, row[i] if i < len(row) else "")
            for run in para.runs:
                run.font.size = Pt(8.5)
            if idx % 2 == 1:
                shade(cell, BAND_FILL)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# --------------------------------------------------------------------------
# Conversion principale
# --------------------------------------------------------------------------
def convert(md_path: str, out_path: str, title: str, subtitle: str) -> None:
    lines = open(md_path, encoding="utf-8").read().split("\n")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    usable = section.page_width - section.left_margin - section.right_margin

    setup_styles(doc)
    cover_page(doc, title, subtitle)
    toc_field(doc)
    page_numbers(doc)

    i = 0
    n = len(lines)
    in_code = False
    code_buffer = []
    skip_toc_block = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Blocs de code
        if stripped.startswith("```"):
            if in_code:
                para = doc.add_paragraph(style="CodeBlock")
                para.add_run("\n".join(code_buffer))
                shade_para(para)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # Le sommaire rédigé à la main est remplacé par le champ TOC de Word.
        if stripped == "# Sommaire":
            skip_toc_block = True
            i += 1
            continue
        if skip_toc_block:
            if stripped.startswith("---"):
                skip_toc_block = False
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Règle horizontale
        if re.fullmatch(r"-{3,}", stripped):
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(4)
            pf.space_after = Pt(8)
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D6DBE1")
            pbdr.append(bottom)
            para._p.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        # Titres
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            para = doc.add_paragraph(style=f"Heading {level}")
            add_inline(para, text)
            i += 1
            continue

        # Tableaux
        if stripped.startswith("|") and i + 1 < n and re.match(
            r"^\s*\|?[\s:|-]+\|[\s:|-]*$", lines[i + 1]
        ):
            rows = [split_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, usable)
            continue

        # Citations
        if stripped.startswith(">"):
            buffer = []
            while i < n and lines[i].strip().startswith(">"):
                buffer.append(lines[i].strip().lstrip(">").strip())
                i += 1
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Inches(0.25)
            pf.space_before = Pt(6)
            pf.space_after = Pt(8)
            pbdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "0F3D63")
            pbdr.append(left)
            para._p.get_or_add_pPr().append(pbdr)
            add_inline(para, " ".join(x for x in buffer if x))
            for run in para.runs:
                run.font.size = Pt(9.5)
            i += 1 if i < n else 0
            continue

        # Listes
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2)
            if len(m.group(1)) >= 2:
                para.paragraph_format.left_indent = Inches(0.6)
            add_inline(para, m.group(2))
            i += 1
            continue

        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            para = doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, m.group(2))
            i += 1
            continue

        # Paragraphe : agrège les lignes consécutives.
        buffer = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith(("#", "|", ">", "```", "- ", "* "))
                or re.match(r"^\d+\.\s", nxt)
                or re.fullmatch(r"-{3,}", nxt)
            ):
                break
            buffer.append(nxt)
            i += 1
        add_inline(doc.add_paragraph(), " ".join(buffer))

    doc.save(out_path)
    print(f"Document généré : {out_path}")


def shade_para(para) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), "F4F6F8")
    para._p.get_or_add_pPr().append(el)


if __name__ == "__main__":
    if len(sys.argv) < 5:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
